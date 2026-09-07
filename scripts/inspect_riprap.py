import sys
import os
import openpyxl
import json

sys.stdout.reconfigure(encoding='utf-8')

base_dir = r"C:\Users\USER\WORKSPACE\00_INBOX\000\TEAM\Doc no. 1211\_Analysis"
wb_master = os.path.join(base_dir, "KZDR_S03_Drainage_Master_Dataset.xlsx")

wb = openpyxl.load_workbook(wb_master, data_only=True)
ws = wb["Slope Protection"]

print(f"Max row: {ws.max_row}")
headers = [cell.value for cell in ws[3]]
print("Headers:", headers)

rows = []
for r in range(4, ws.max_row + 1):
    row_vals = [cell.value for cell in ws[r]]
    if any(row_vals):
        rows.append(row_vals)

print(f"Total non-empty data rows: {len(rows)}")
for i in [0, 1, 2, 50, 100, 200, -2, -1]:
    if 0 <= i < len(rows) or (i < 0 and abs(i) <= len(rows)):
        print(f"Row sample {i}: {rows[i]}")
