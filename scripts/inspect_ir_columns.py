import openpyxl

wb_path = r"C:\Users\USER\Downloads\IR Registers\WSO-REG-COR-KMD-02-IRS-00 - Inspection Requests IN-OUT (S02,S03).xlsx"
wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
sheet = wb["IR (0001-1000)"]

print("Rows 8 to 15:")
row_idx = 0
for row in sheet.iter_rows(values_only=True):
    row_idx += 1
    if 8 <= row_idx <= 15:
        # filter out trailing Nones
        last_non_none = -1
        for i, v in enumerate(row):
            if v is not None:
                last_non_none = i
        print(f"\n--- Row {row_idx} ---")
        for i in range(last_non_none + 1):
            if row[i] is not None:
                print(f"Col {i} ({openpyxl.utils.get_column_letter(i+1)}): {row[i]}")

wb.close()
