import openpyxl
import os

files = [
    r"Doc no. 1211/_Analysis/KZDR_App_Correction_Register.xlsx",
    r"KZDR_Section03_Drainage_IR_Map.xlsx"
]

for fp in files:
    if not os.path.exists(fp):
        print(f"File not found: {fp}")
        continue
    wb = openpyxl.load_workbook(fp, read_only=True)
    print(f"\n=== {fp} ===")
    print("Sheets:", wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(max_row=5, values_only=True))
    for r in rows:
        print(" ", r[:10])
    wb.close()
