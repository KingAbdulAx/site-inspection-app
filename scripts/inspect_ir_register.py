import openpyxl
import os

wb_path = r"C:\Users\USER\Downloads\IR Registers\WSO-REG-COR-KMD-02-IRS-00 - Inspection Requests IN-OUT (S02,S03).xlsx"
print(f"Loading workbook: {wb_path}")

wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
print("Sheet names:")
for name in wb.sheetnames:
    print(f"  - {name}")

for name in wb.sheetnames:
    sheet = wb[name]
    print(f"\n--- Sheet: {name} (sample rows) ---")
    row_count = 0
    for row in sheet.iter_rows(values_only=True):
        row_count += 1
        non_empty = [c for c in row if c is not None]
        if non_empty:
            print(f"Row {row_count}: {row[:12]}")
        if row_count >= 10:
            break

wb.close()
