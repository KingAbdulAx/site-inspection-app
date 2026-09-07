import sys
import os
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

base_dir = r"C:\Users\USER\WORKSPACE\00_INBOX\000\TEAM\Doc no. 1211\_Analysis"
wb_master = os.path.join(base_dir, "KZDR_S03_Drainage_Master_Dataset.xlsx")
wb_corr = os.path.join(base_dir, "KZDR_App_Correction_Register.xlsx")

print("Checking Master Dataset sheets:")
wb1 = openpyxl.load_workbook(wb_master, data_only=True)
print("Sheets in master:", wb1.sheetnames)

if "Slope Protection" in wb1.sheetnames:
    ws_sp = wb1["Slope Protection"]
    print(f"Slope Protection sheet max_row={ws_sp.max_row}, max_column={ws_sp.max_column}")
    print("Row 1:", [cell.value for cell in ws_sp[1]])
    print("Row 2:", [cell.value for cell in ws_sp[2]])
    print("Row 3:", [cell.value for cell in ws_sp[3]])
    print("Row 4 sample:", [cell.value for cell in ws_sp[4]])
    print("Row 5 sample:", [cell.value for cell in ws_sp[5]])

if "Water Descents" in wb1.sheetnames:
    ws_wd = wb1["Water Descents"]
    print("Water Descents in master:", ws_wd.max_row)
    print("Row 1:", [cell.value for cell in ws_wd[1]])
    print("Row 2 sample:", [cell.value for cell in ws_wd[2]])

print("\nChecking Correction Register sheets:")
wb2 = openpyxl.load_workbook(wb_corr, data_only=True)
print("Sheets in corr:", wb2.sheetnames)

if "Missing Classes" in wb2.sheetnames:
    ws_mc = wb2["Missing Classes"]
    print(f"Missing Classes max_row={ws_mc.max_row}, max_column={ws_mc.max_column}")
    for r in range(1, min(25, ws_mc.max_row + 1)):
        print(f"Row {r}:", [cell.value for cell in ws_mc[r]])
