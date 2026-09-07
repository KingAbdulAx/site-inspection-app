import sys
import os
import openpyxl
import json
import re
import math

sys.stdout.reconfigure(encoding='utf-8')

# 1. Test Centerline
with open('data/section03_centerline.json', 'r', encoding='utf-8') as f:
    centerline_data = json.load(f)

dense_points = centerline_data['dense_points']
pks = [p['pk'] for p in dense_points]
print(f"Centerline loaded: {len(dense_points)} points, PK {pks[0]:.1f} to {pks[-1]:.1f}")

def get_track_point(target_pk):
    if target_pk <= pks[0]:
        return dense_points[0]['lon'], dense_points[0]['lat'], dense_points[0]['bearing']
    if target_pk >= pks[-1]:
        return dense_points[-1]['lon'], dense_points[-1]['lat'], dense_points[-1]['bearing']
    idx = 0
    while idx < len(pks) - 1 and pks[idx+1] < target_pk:
        idx += 1
    p1 = dense_points[idx]
    p2 = dense_points[idx+1]
    denom = p2['pk'] - p1['pk']
    frac = (target_pk - p1['pk']) / denom if denom > 0 else 0
    t_lon = p1['lon'] + frac * (p2['lon'] - p1['lon'])
    t_lat = p1['lat'] + frac * (p2['lat'] - p1['lat'])
    return t_lon, t_lat, p1['bearing']

def offset_point(lon, lat, bearing, dist_m):
    R = 6371000.0
    offset_bearing = (bearing + (90 if dist_m >= 0 else -90)) % 360
    d = abs(dist_m)
    phi1 = math.radians(lat)
    lam1 = math.radians(lon)
    brg = math.radians(offset_bearing)
    delta = d / R
    phi2 = math.asin(math.sin(phi1) * math.cos(delta) + math.cos(phi1) * math.sin(delta) * math.cos(brg))
    lam2 = lam1 + math.atan2(math.sin(brg) * math.sin(delta) * math.cos(phi1), math.cos(delta) - math.sin(phi1) * math.sin(phi2))
    return math.degrees(lam2), math.degrees(phi2)

# 2. Test Crest Descent at 84+544
t_lon, t_lat, brg = get_track_point(84544.0)
c_lon, c_lat = offset_point(t_lon, t_lat, brg, -22.0)
print(f"Crest descent coord at PK 84+544: lon={c_lon:.6f}, lat={c_lat:.6f}")

# 3. Test Riprap
base_dir = r"..\Doc no. 1211\_Analysis"
wb_master = os.path.join(base_dir, "KZDR_S03_Drainage_Master_Dataset.xlsx")
wb = openpyxl.load_workbook(wb_master, data_only=True)
ws_sp = wb["Slope Protection"]

def parse_chainage(ch_str):
    if isinstance(ch_str, (int, float)):
        return float(ch_str), f"PK {int(ch_str//1000)}+{int(ch_str%1000):03d}"
    m = re.search(r'(\d+)\+(\d+)(?:\.(\d+))?', str(ch_str))
    if not m:
        raise ValueError(f"Cannot parse chainage: {ch_str}")
    km = int(m.group(1))
    meters = float(m.group(2))
    if m.group(3):
        meters += float('.' + m.group(3))
    pk = km * 1000.0 + meters
    m_int = int(round(meters))
    ch_formatted = f"PK {km}+{m_int:03d}"
    return pk, ch_formatted

riprap_items = []
for r in range(4, ws_sp.max_row + 1):
    vals = [ws_sp.cell(r, c).value for c in range(1, 9)]
    if not any(vals):
        continue
    sp_id, side, ch_raw, l_val, h_val, d50_val, sheet, callout = vals
    pk, ch_str = parse_chainage(ch_raw)
    side_str = "Right" if side == 'R' else "Left"
    sign = 1.0 if side == 'R' else -1.0
    t_lon, t_lat, brg = get_track_point(pk)
    pt_lon, pt_lat = offset_point(t_lon, t_lat, brg, sign * 18.0)
    riprap_items.append({
        'id': sp_id,
        'pk': pk,
        'ch_str': ch_str,
        'side': side_str,
        'sheet': sheet,
        'callout': callout,
        'coord': [pt_lon, pt_lat]
    })

print(f"Total riprap parsed: {len(riprap_items)}, first PK {riprap_items[0]['pk']}, last PK {riprap_items[-1]['pk']}")

# 4. Test Water Descents from KZDR_Field.html
html_path = r"..\kzdr-field\KZDR_Field.html"
with open(html_path, "r", encoding="utf-8") as f:
    text = f.read()

m = re.search(r"const FEATURES\s*=\s*(\{.*?\});", text, re.DOTALL)
data = json.loads(m.group(1))
wd_raw = [f for f in data.get('features', []) if f.get('type') == 'water_descent']
print(f"Total water descents in KZDR_Field.html: {len(wd_raw)}")

wd_items = []
for w in wd_raw:
    pk = float(w['point_chainage'])
    km = int(pk // 1000)
    m_val = pk % 1000
    ch_str = f"PK {km}+{int(round(m_val)):03d}"
    side_str = "Right" if w.get('side') == 'R' else "Left"
    sign = 1.0 if w.get('side') == 'R' else -1.0
    offset = float(w.get('offset_m') or 15.0)
    t_lon, t_lat, brg = get_track_point(pk)
    pt_lon, pt_lat = offset_point(t_lon, t_lat, brg, sign * offset)
    wd_items.append({
        'id': w['id'],
        'pk': pk,
        'ch_str': ch_str,
        'side': side_str,
        'offset': offset,
        'sheet': w.get('refs', {}).get('drawings', ['DW-03000'])[0],
        'coord': [pt_lon, pt_lat]
    })

print(f"Total water descents parsed: {len(wd_items)}, first PK {wd_items[0]['pk']}, last PK {wd_items[-1]['pk']}")
