import openpyxl
import os

excel_path = os.path.abspath(r"Doc no. 1211\_Analysis\KZDR_S03_Drainage_Master_Dataset.xlsx")
wb = openpyxl.load_workbook(excel_path, data_only=True)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    for row in ws.iter_rows(values_only=True):
        row_str = " | ".join([str(c) for c in row if c is not None])
        if '88.3' in row_str or '88+27' in row_str or '88+3' in row_str or 'bc88' in row_str.lower() or 'bc 88' in row_str.lower():
            print(f"[{sheetname}] {row_str}")
