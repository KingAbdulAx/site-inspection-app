import json
import math
import openpyxl
import re
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

# 1. Load Centerline
with open('data/section03_centerline.json', 'r', encoding='utf-8') as f:
    centerline_data = json.load(f)

dense_points = centerline_data['dense_points']
pks = [p['pk'] for p in dense_points]

def get_track_point(target_pk):
    if target_pk <= pks[0]:
        return dense_points[0]['lon'], dense_points[0]['lat'], dense_points[0]['bearing']
    if target_pk >= pks[-1]:
        return dense_points[-1]['lon'], dense_points[-1]['lat'], dense_points[-1]['bearing']
    idx = 0
    while idx < len(pks) - 1 and pks[idx+1] < target_pk:
        idx += 1
    p1 = dense_points[idx]
    p2 = dense_points[idx+1]
    denom = p2['pk'] - p1['pk']
    frac = (target_pk - p1['pk']) / denom if denom > 0 else 0
    t_lon = p1['lon'] + frac * (p2['lon'] - p1['lon'])
    t_lat = p1['lat'] + frac * (p2['lat'] - p1['lat'])
    return t_lon, t_lat, p1['bearing']

def offset_point(lon, lat, bearing, dist_m):
    R = 6371000.0
    offset_bearing = (bearing + (90 if dist_m >= 0 else -90)) % 360
    d = abs(dist_m)
    phi1 = math.radians(lat)
    lam1 = math.radians(lon)
    brg = math.radians(offset_bearing)
    delta = d / R
    phi2 = math.asin(math.sin(phi1) * math.cos(delta) + math.cos(phi1) * math.sin(delta) * math.cos(brg))
    lam2 = lam1 + math.atan2(math.sin(brg) * math.sin(delta) * math.cos(phi1), math.cos(delta) - math.sin(phi1) * math.sin(phi2))
    return math.degrees(lam2), math.degrees(phi2)

def parse_chainage(ch_str):
    if isinstance(ch_str, (int, float)):
        return float(ch_str), f"PK {int(ch_str//1000)}+{int(ch_str%1000):03d}"
    m = re.search(r'(\d+)\+(\d+)(?:\.(\d+))?', str(ch_str))
    if not m:
        raise ValueError(f"Cannot parse chainage: {ch_str}")
    km = int(m.group(1))
    meters = float(m.group(2))
    if m.group(3):
        meters += float('.' + m.group(3))
    pk = km * 1000.0 + meters
    m_int = int(round(meters))
    ch_formatted = f"PK {km}+{m_int:03d}"
    return pk, ch_formatted

# 2. Load existing assets and snapshot
with open('data/section03_assets.json', 'r', encoding='utf-8') as f:
    assets_data = json.load(f)

features = assets_data['features']
print(f"Starting assets count: {len(features)}")

with open('data/field_data_snapshot.json', 'r', encoding='utf-8') as f:
    snapshot = json.load(f)

# Verify initial match of baseline features
for snap_id, snap_data in snapshot.items():
    feat = next((f for f in features if f['properties']['id'] == snap_id), None)
    assert feat is not None, f"Baseline feature {snap_id} missing!"
    p = feat['properties']
    assert p.get('status') == snap_data['status'], f"Status mismatch on {snap_id}"
    assert p.get('inspection_date') == snap_data['inspection_date'], f"Date mismatch on {snap_id}"
    if snap_data['notes']:
        assert snap_data['notes'] in p.get('notes', ''), f"Original notes missing on {snap_id}"

print("Field data continuity check before Phase 6 additions: 100% OK")

new_features = []
changelog_lines = []

# --- ITEM 1: Crest Ditch Descent at PK 84+544 (1 feature) ---
t_lon, t_lat, brg = get_track_point(84544.0)
p_lon, p_lat = offset_point(t_lon, t_lat, brg, -22.0)

feat_descent = {
    "type": "Feature",
    "id": "asset_265",
    "properties": {
        "id": "asset_265",
        "chainage_str": "PK 84+544",
        "start_pk": 84544.0,
        "end_pk": 84544.0,
        "length_m": 0.0,
        "is_point": True,
        "side": "Left",
        "position": "Left Cut Crest Descent",
        "typology": "Concrete Crest Ditch Descent (to Type 1 Side Ditch)",
        "short_code": "Descent",
        "category": "Energy Dissipator",
        "color": "#D97706",
        "drawing_ref": "DW-03002-07 / DW-10003-04-A",
        "specs": "Concrete crest ditch descent into Type 1 side ditch (i=2.25%, fall=0.24m) per DW-03002-07 and standard detail DW-10003-04-A.",
        "icon": "cascade",
        "status": "Not Started",
        "notes": "[EVIDENCE: VERIFIED] Discrete drop descent connecting Type 11 crest ditch to Type 1 platform ditch at CH 84+544.",
        "inspection_date": ""
    },
    "geometry": {
        "type": "Point",
        "coordinates": [round(p_lon, 7), round(p_lat, 7)]
    }
}
new_features.append(feat_descent)
changelog_lines.append("| `asset_265` | `PK 84+544` | `ADD (Point)` | Crest Ditch Descent | DW-03002-07 / DW-10003-04-A | Concrete crest ditch descent (i=2.25%, fall=0.24m) into Type 1 side ditch |")

# --- ITEM 2: Riprap Slope Protection (267 features) ---
base_dir = r"..\Doc no. 1211\_Analysis"
wb_master = os.path.join(base_dir, "KZDR_S03_Drainage_Master_Dataset.xlsx")
wb = openpyxl.load_workbook(wb_master, data_only=True)
ws_sp = wb["Slope Protection"]

next_asset_id = 266
for r in range(4, ws_sp.max_row + 1):
    vals = [ws_sp.cell(r, c).value for c in range(1, 9)]
    if not any(vals):
        continue
    sp_id, side, ch_raw, l_val, h_val, d50_val, sheet, callout = vals
    pk, ch_str = parse_chainage(ch_raw)
    side_str = "Right" if side == 'R' else "Left"
    sign = 1.0 if side == 'R' else -1.0
    
    t_lon, t_lat, brg = get_track_point(pk)
    pt_lon, pt_lat = offset_point(t_lon, t_lat, brg, sign * 18.0)
    
    length_m = float(l_val) if l_val and str(l_val).replace('.', '', 1).isdigit() else 0.0
    
    if callout and str(callout).strip():
        specs_str = f"{str(callout).strip()}. Standard detail DW-10004-02-A."
    else:
        specs_str = "Rock riprap embankment protection armor per standard detail DW-10004-02-A."
        
    feat_id = f"asset_{next_asset_id:03d}"
    feat_riprap = {
        "type": "Feature",
        "id": feat_id,
        "properties": {
            "id": feat_id,
            "ref_id": sp_id,
            "chainage_str": ch_str,
            "start_pk": pk,
            "end_pk": pk,
            "length_m": length_m,
            "is_point": True,
            "side": side_str,
            "position": "Embankment Slope Toe (Riprap Armor)",
            "typology": "Embankment Riprap Armor & Scour Protection",
            "short_code": "Riprap",
            "category": "Riprap Protection",
            "color": "#78716C",
            "drawing_ref": f"{sheet} / DW-10004-02-A",
            "specs": specs_str,
            "icon": "protection",
            "status": "Not Started",
            "notes": "[DESIGN NOTE] Position is the AutoCAD callout label anchor, not the physical start/end extent of the rock armor. Marks stretch where longitudinal ditch is omitted.",
            "inspection_date": ""
        },
        "geometry": {
            "type": "Point",
            "coordinates": [round(pt_lon, 7), round(pt_lat, 7)]
        }
    }
    new_features.append(feat_riprap)
    next_asset_id += 1

print(f"Added riprap features: {len(new_features) - 1} (asset_266 to asset_{next_asset_id - 1:03d})")
changelog_lines.append(f"| `asset_266`..`asset_532` (267 features) | `PK 82+939` to `PK 124+485` | `ADD (Point)` | Slope Protection (SP-001 to SP-267) | DW-03001 to DW-03030 / DW-10004-02-A | Embankment riprap armor callouts marking stretches where longitudinal ditch is omitted |")

# --- ITEM 3: Water Descents (310 features) ---
html_path = r"..\kzdr-field\KZDR_Field.html"
with open(html_path, "r", encoding="utf-8") as f:
    text = f.read()

m = re.search(r"const FEATURES\s*=\s*(\{.*?\});", text, re.DOTALL)
data = json.loads(m.group(1))
wd_raw = [f for f in data.get('features', []) if f.get('type') == 'water_descent']
print(f"Total water descents loaded from KZDR_Field.html: {len(wd_raw)}")

wd_start_id = next_asset_id
for w in wd_raw:
    pk = float(w['point_chainage'])
    km = int(pk // 1000)
    m_val = pk % 1000
    ch_str = f"PK {km}+{int(round(m_val)):03d}"
    side_str = "Right" if w.get('side') == 'R' else "Left"
    sign = 1.0 if w.get('side') == 'R' else -1.0
    offset = float(w.get('offset_m') or 15.0)
    
    t_lon, t_lat, brg = get_track_point(pk)
    pt_lon, pt_lat = offset_point(t_lon, t_lat, brg, sign * offset)
    
    sheet_ref = w.get('refs', {}).get('drawings', ['DW-03000'])[0]
    
    feat_id = f"asset_{next_asset_id:03d}"
    feat_wd = {
        "type": "Feature",
        "id": feat_id,
        "properties": {
            "id": feat_id,
            "ref_id": w['id'],
            "chainage_str": ch_str,
            "start_pk": pk,
            "end_pk": pk,
            "length_m": 0.0,
            "is_point": True,
            "side": side_str,
            "position": "Embankment Slope Face (Shoulder Chute)",
            "typology": "Precast Water Descent (Type 9 Chute)",
            "short_code": "Desc",
            "category": "Water Descent",
            "color": "#0284C7",
            "drawing_ref": f"{sheet_ref} / DW-10003-04-A",
            "specs": "Precast half-round chute D=0.30m down embankment slope discharging shoulder ditch into toe dissipator per DW-10003-04-A.",
            "icon": "cascade",
            "status": "Not Started",
            "notes": "[EVIDENCE: PROBABLE] Extracted from solid-blue symbol pair on plan sheet. Median spacing ~44m; resolves hydraulic length limit of Type 9 shoulder ditch.",
            "inspection_date": ""
        },
        "geometry": {
            "type": "Point",
            "coordinates": [round(pt_lon, 7), round(pt_lat, 7)]
        }
    }
    new_features.append(feat_wd)
    next_asset_id += 1

print(f"Added water descent features: {len(wd_raw)} (asset_{wd_start_id:03d} to asset_{next_asset_id - 1:03d})")
changelog_lines.append(f"| `asset_533`..`asset_842` (310 features) | `PK 82+911` to `PK 124+459` | `ADD (Point)` | Water Descents (WD-001 to WD-310) | DW-03001 to DW-03030 / DW-10003-04-A | Precast half-round chutes preventing shoulder ditches exceeding permissible hydraulic run lengths |")

# Append all new features
all_features = features + new_features
print(f"Final total features count: {len(all_features)}")

# Update assets_data metadata
assets_data['features'] = all_features
assets_data['metadata']['count'] = len(all_features)

# Verify Sacred Field Data Preservation
print("\n--- Running Automated Sacred Field Data Diff ---")
lost_status = 0
lost_notes = 0
lost_dates = 0

for snap_id, snap_data in snapshot.items():
    feat = next((f for f in all_features if f['properties']['id'] == snap_id), None)
    if not feat:
        print(f"CRITICAL ERROR: Feature {snap_id} missing from final dataset!")
        sys.exit(1)
    p = feat['properties']
    if p.get('status') != snap_data['status']:
        print(f"ERROR: Status altered on {snap_id}: '{p.get('status')}' != '{snap_data['status']}'")
        lost_status += 1
    if snap_data['notes'] and snap_data['notes'] not in p.get('notes', ''):
        print(f"ERROR: Notes lost on {snap_id}: '{snap_data['notes']}' not in '{p.get('notes')}'")
        lost_notes += 1
    if p.get('inspection_date') != snap_data['inspection_date']:
        print(f"ERROR: Date altered on {snap_id}: '{p.get('inspection_date')}' != '{snap_data['inspection_date']}'")
        lost_dates += 1

print(f"Lost status count: {lost_status}")
print(f"Lost notes count: {lost_notes}")
print(f"Lost inspection dates: {lost_dates}")

if lost_status == 0 and lost_notes == 0 and lost_dates == 0:
    print("SUCCESS: 100% field data continuity verified. Writing output file...")
    with open('data/section03_assets.json', 'w', encoding='utf-8') as f:
        json.dump(assets_data, f, indent=2, ensure_ascii=False)
    print("data/section03_assets.json written successfully.")
else:
    print("ABORTING: Field data preservation check failed!")
    sys.exit(1)

# Append to CORRECTIONS_CHANGELOG.md
changelog_path = 'CORRECTIONS_CHANGELOG.md'
with open(changelog_path, 'a', encoding='utf-8') as f:
    f.write('\n### Phase 6 — Missing Feature Classes & Crest Descent Point Feature\n\n')
    f.write('| Record ID | Chainage Extent | Action | Register Reference | Drawing / Document Reference | Notes & Field Data Continuity |\n')
    f.write('|:---|:---|:---|:---|:---|:---|\n')
    for line in changelog_lines:
        f.write(line + '\n')

print("CORRECTIONS_CHANGELOG.md updated successfully.")
