import openpyxl

wb = openpyxl.load_workbook("app/data/KMD_Drainage_Extraction_Batch_Verification_Workbook.xlsx", data_only=False)
print("Sheet names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\nSheet: {sname}, dimensions: {ws.dimensions}")
    headers = [cell.value for cell in ws[1]] if sname != "Batch Review Dashboard" else [cell.value for cell in ws[6]]
    print("  Headers:", headers[:8], "...")
    if sname == "Batch Review Dashboard":
        print(f"  Total batches listed: {ws.max_row - 6}")
    else:
        print(f"  Total feature rows: {ws.max_row - 1}")
