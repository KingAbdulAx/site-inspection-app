import openpyxl
import json
import math
import re

# Load centerline
with open('data/section03_centerline.json', 'r', encoding='utf-8') as f:
    centerline_data = json.load(f)

dense_points = centerline_data['dense_points']
pks = [p['pk'] for p in dense_points]

def get_track_point(target_pk):
    if target_pk <= pks[0]:
        return dense_points[0]['lon'], dense_points[0]['lat'], dense_points[0]['bearing']
    if target_pk >= pks[-1]:
        return dense_points[-1]['lon'], dense_points[-1]['lat'], dense_points[-1]['bearing']
    idx = 0
    while idx < len(pks) - 1 and pks[idx+1] < target_pk:
        idx += 1
    p1 = dense_points[idx]
    p2 = dense_points[idx+1]
    denom = p2['pk'] - p1['pk']
    frac = (target_pk - p1['pk']) / denom if denom > 0 else 0
    t_lon = p1['lon'] + frac * (p2['lon'] - p1['lon'])
    t_lat = p1['lat'] + frac * (p2['lat'] - p1['lat'])
    return t_lon, t_lat, p1['bearing']

def offset_point(lon, lat, bearing, dist_m):
    R = 6371000.0
    offset_bearing = (bearing + (90 if dist_m >= 0 else -90)) % 360
    d = abs(dist_m)
    phi1 = math.radians(lat)
    lam1 = math.radians(lon)
    brg = math.radians(offset_bearing)
    delta = d / R
    phi2 = math.asin(math.sin(phi1) * math.cos(delta) + math.cos(phi1) * math.sin(delta) * math.cos(brg))
    lam2 = lam1 + math.atan2(math.sin(brg) * math.sin(delta) * math.cos(phi1), math.cos(delta) - math.sin(phi1) * math.sin(phi2))
    return math.degrees(lam2), math.degrees(phi2)

def parse_ditch_chainage(ch_str):
    m = re.search(r'(\d+)\+(\d+)', ch_str)
    if not m:
        raise ValueError(f'Cannot parse chainage: {ch_str}')
    km = int(m.group(1))
    meters = float(m.group(2))
    pk = km * 1000.0 + meters
    m_int = int(round(meters))
    ch_formatted = f'PK {km}+{m_int:03d}'
    return pk, ch_formatted

# Load target assets
with open('data/section03_assets.json', 'r', encoding='utf-8') as f:
    assets_data = json.load(f)

features = assets_data['features']
print(f'Starting features count in Phase 3: {len(features)}')

# Load snapshot
with open('data/field_data_snapshot.json', 'r', encoding='utf-8') as f:
    snapshot = json.load(f)

# Load register
wb = openpyxl.load_workbook(r'..\Doc no. 1211\_Analysis\KZDR_App_Correction_Register.xlsx', data_only=True)
sheet = wb['Ditch Runs']

added_features = []
changelog_entries = []

next_asset_num = len(features) + 1

# Process ADD rows
for r in range(4, sheet.max_row+1):
    vals = [sheet.cell(r, c).value for c in range(1, 10)]
    act = vals[0]
    if act == 'ADD':
        ch_from = vals[1]
        ch_to = vals[2]
        length_val = float(vals[3])
        side_code = str(vals[4]).strip()
        types_str = str(vals[5]).strip()
        ref_ids = str(vals[6]).strip()
        sheet_ref = str(vals[7]).strip()
        note_str = str(vals[8]).strip() if vals[8] else ''

        pk_start, ch_from_fmt = parse_ditch_chainage(ch_from)
        pk_end, ch_to_fmt = parse_ditch_chainage(ch_to)
        
        # Siding and offset
        if side_code == 'L':
            side = 'Left'
            sign = -1
        else:
            side = 'Right'
            sign = 1
            
        if 'Type 1' in types_str and 'Type 12' not in types_str and 'Type 7' not in types_str and 'Type 4' not in types_str:
            offset_dist = 5.5 * sign
            category = 'Side Ditch'
            color = '#3B82F6'
            position = f'{side} Side (Track Cut / Platform)'
            typology = 'Concrete Lined Side Ditch (Type 1 Standard)'
            short_code = 'Type 1 Std'
            specs = f'Concrete trapezoidal cut side ditch (b=0.75m, h=0.75m, 1:1) per {sheet_ref} [Ref: {ref_ids}]'
        elif 'Type 4' in types_str and 'Type 12' not in types_str and 'Type 7' not in types_str:
            offset_dist = 18.0 * sign
            category = 'Toe Ditch'
            color = '#92400E'
            position = f'{side} Side (Embankment Toe)'
            typology = 'Unlined Earth Toe Ditch (Type 4)'
            short_code = 'Type 4'
            specs = f'Unlined earth ditch at embankment foot (H=1.0m, 1:1.5) per {sheet_ref} [Ref: {ref_ids}]'
        elif 'Type 7' in types_str and 'Type 12' not in types_str and 'Type 4' not in types_str:
            offset_dist = 17.0 * sign
            category = 'Toe Ditch'
            color = '#EA580C'
            position = f'{side} Side (Embankment Toe)'
            typology = 'Concrete Lined Toe Ditch (Type 7)'
            short_code = 'Type 7'
            specs = f'Concrete lined triangular toe ditch (H=0.80m, 1:1.5) per {sheet_ref} [Ref: {ref_ids}]'
        elif 'Type 12' in types_str and 'Type 7' not in types_str and 'Type 4' not in types_str:
            offset_dist = 18.0 * sign
            category = 'Toe Ditch'
            color = '#F59E0B'
            position = f'{side} Side (Embankment Toe)'
            typology = 'Trapezoidal Lined Toe Ditch (Type 12)'
            short_code = 'Type 12'
            specs = f'Concrete trapezoidal toe ditch (B=1.50m, H=0.50m, 1:1.5) per {sheet_ref} [Ref: {ref_ids}]'
        else:
            offset_dist = 17.5 * sign
            category = 'Toe Ditch'
            color = '#F59E0B'
            position = f'{side} Side (Embankment Toe)'
            typology = f'Concrete Lined Toe Ditch ({types_str})'
            short_code = 'Toe Ditch'
            specs = f'Concrete lined toe ditch ({types_str}) per {sheet_ref} [Ref: {ref_ids}]'

        # Generate LineString coordinates
        coords = []
        span_m = pk_end - pk_start
        steps = max(2, int(math.ceil(abs(span_m) / 25.0)))
        for s in range(steps + 1):
            cur_pk = pk_start + (s / steps) * span_m
            t_lon, t_lat, brg = get_track_point(cur_pk)
            p_off = offset_point(t_lon, t_lat, brg, offset_dist)
            coords.append([round(p_off[0], 7), round(p_off[1], 7)])

        asset_id = f'asset_{next_asset_num:03d}'
        next_asset_num += 1
        
        chainage_str = f'{ch_from_fmt} – {ch_to_fmt}'
        drawing_ref = f'{sheet_ref} / MDDTDW220010001'
        
        feature = {
            'type': 'Feature',
            'id': asset_id,
            'properties': {
                'id': asset_id,
                'chainage_str': chainage_str,
                'start_pk': pk_start,
                'end_pk': pk_end,
                'length_m': round(length_val, 1),
                'is_point': False,
                'side': side,
                'position': position,
                'typology': typology,
                'short_code': short_code,
                'category': category,
                'color': color,
                'drawing_ref': drawing_ref,
                'specs': specs,
                'icon': 'ditch',
                'status': 'Not Started',
                'notes': '',
                'inspection_date': ''
            },
            'geometry': {
                'type': 'LineString',
                'coordinates': coords
            }
        }
        added_features.append(feature)
        changelog_entries.append(
            f'| {asset_id} | {chainage_str} ({side}) | ADD | Ditch Runs (Row {r}) | {drawing_ref} | Added longitudinal ditch run {typology} ({ref_ids}). Initial field data intact. |'
        )

# Process RE-EVIDENCE rows
reev_records_flagged = set()
for r in range(42, 58):
    vals = [sheet.cell(r, c).value for c in range(1, 10)]
    act = vals[0]
    if act == 'RE-EVIDENCE':
        ch_from = vals[1]
        ch_to = vals[2]
        side_code = vals[4]
        types_str = vals[5]
        ref_id = str(vals[6]).strip()
        sheet_ref = str(vals[7]).strip()
        
        # Find matching feature
        for feat in features:
            if feat['properties']['id'] == ref_id:
                reev_records_flagged.add(ref_id)
                existing_notes = feat['properties'].get('notes', '')
                if ref_id in ['asset_090', 'asset_091']:
                    flag_note = f'[FLAG] No drawing symbology found on {sheet_ref}; alignment intersects railway bridge BRG-2602 at CH 93+678–93+731 where continuous platform side ditch is structurally precluded. Marked for Senior Engineer review.'
                else:
                    flag_note = f'[FLAG] No drawing symbology found on {sheet_ref} supporting this run. Retained pending site verification.'
                
                if flag_note not in existing_notes:
                    feat['properties']['notes'] = (existing_notes + ' ' + flag_note).strip()
                
                ch_str = feat['properties']['chainage_str']
                changelog_entries.append(
                    f'| {ref_id} | {ch_str} | FLAG | Ditch Runs (Row {r}) | {sheet_ref} | {flag_note} Existing record retained. |'
                )

# Append added features
features.extend(added_features)
assets_data['metadata']['count'] = len(features)

print(f'Total features after Phase 3: {len(features)} (Added {len(added_features)} ditch runs, Flagged {len(reev_records_flagged)} re-evidence assets)')

# Diff against snapshot
for fid, s_data in snapshot.items():
    matching = [f for f in features if f['properties']['id'] == fid]
    assert len(matching) == 1, f'Missing feature {fid}'
    f_props = matching[0]['properties']
    # status must match exactly
    assert f_props['status'] == s_data['status'], f'Status modified for {fid}'
    # if fid not in flagged items, notes must match snapshot exactly
    if fid not in ['asset_094', 'asset_133'] and fid not in reev_records_flagged:
        assert f_props['notes'] == s_data['notes'], f'Notes modified for {fid}'
    assert f_props['inspection_date'] == s_data['inspection_date'], f'Inspection date modified for {fid}'

print('Field data verification passed! 0 records lost or overwritten.')

# Write back section03_assets.json
with open('data/section03_assets.json', 'w', encoding='utf-8') as f:
    json.dump(assets_data, f, indent=2, ensure_ascii=False)

# Append to CORRECTIONS_CHANGELOG.md
with open('CORRECTIONS_CHANGELOG.md', 'a', encoding='utf-8') as f:
    f.write('\n### Phase 3 — Longitudinal Ditch Runs Corrections (Ditch Runs Tab)\n\n')
    f.write('| Record ID | Chainage Extent | Action | Register Reference | Drawing / Document Reference | Notes & Field Data Continuity |\n')
    f.write('|:---|:---|:---|:---|:---|:---|\n')
    for line in changelog_entries:
        f.write(line + '\n')

print('CORRECTIONS_CHANGELOG.md updated successfully.')
