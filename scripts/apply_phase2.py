import openpyxl
import json
import math
import re

# Load centerline
with open('data/section03_centerline.json', 'r', encoding='utf-8') as f:
    centerline_data = json.load(f)

dense_points = centerline_data['dense_points']
pks = [p['pk'] for p in dense_points]

def get_track_point(target_pk):
    if target_pk <= pks[0]:
        return dense_points[0]['lon'], dense_points[0]['lat'], dense_points[0]['bearing']
    if target_pk >= pks[-1]:
        return dense_points[-1]['lon'], dense_points[-1]['lat'], dense_points[-1]['bearing']
    idx = 0
    while idx < len(pks) - 1 and pks[idx+1] < target_pk:
        idx += 1
    p1 = dense_points[idx]
    p2 = dense_points[idx+1]
    denom = p2['pk'] - p1['pk']
    frac = (target_pk - p1['pk']) / denom if denom > 0 else 0
    t_lon = p1['lon'] + frac * (p2['lon'] - p1['lon'])
    t_lat = p1['lat'] + frac * (p2['lat'] - p1['lat'])
    return t_lon, t_lat, p1['bearing']

def offset_point(lon, lat, bearing, dist_m):
    R = 6371000.0
    offset_bearing = (bearing + (90 if dist_m >= 0 else -90)) % 360
    d = abs(dist_m)
    phi1 = math.radians(lat)
    lam1 = math.radians(lon)
    brg = math.radians(offset_bearing)
    delta = d / R
    phi2 = math.asin(math.sin(phi1) * math.cos(delta) + math.cos(phi1) * math.sin(delta) * math.cos(brg))
    lam2 = lam1 + math.atan2(math.sin(brg) * math.sin(delta) * math.cos(phi1), math.cos(delta) - math.sin(phi1) * math.sin(phi2))
    return math.degrees(lam2), math.degrees(phi2)

def parse_culvert_chainage(ch_str):
    m = re.search(r'(\d+)\+(\d+)(?:\.(\d+))?', ch_str)
    if not m:
        raise ValueError(f'Cannot parse chainage: {ch_str}')
    km = int(m.group(1))
    meters = float(m.group(2))
    if m.group(3):
        meters += float('.' + m.group(3))
    pk = km * 1000.0 + meters
    m_int = int(round(meters))
    ch_formatted = f'PK {km}+{m_int:03d}'
    return pk, ch_formatted

# Load target assets
with open('data/section03_assets.json', 'r', encoding='utf-8') as f:
    assets_data = json.load(f)

features = assets_data['features']
print(f'Starting features count: {len(features)}')

# Load snapshot
with open('data/field_data_snapshot.json', 'r', encoding='utf-8') as f:
    snapshot = json.load(f)

# Verify initial match
assert len(features) == len(snapshot), 'Initial feature count mismatch with snapshot'

# Load register
wb = openpyxl.load_workbook(r'..\\Doc no. 1211\\_Analysis\\KZDR_App_Correction_Register.xlsx', data_only=True)
sheet = wb['Culverts']

added_features = []
changelog_entries = []

next_asset_num = len(features) + 1

for r in range(4, sheet.max_row+1):
    vals = [sheet.cell(r, c).value for c in range(1, 11)]
    act = vals[0]
    if act == 'ADD':
        ch_raw = vals[1]
        cid = vals[2]
        basin = vals[3]
        ctype = vals[4]
        cno = int(vals[5])
        sec = vals[6]
        q = vals[7]
        sheet_ref = vals[8]
        
        pk, ch_formatted = parse_culvert_chainage(ch_raw)
        
        # Geometry
        t_lon, t_lat, brg = get_track_point(pk)
        p_l = offset_point(t_lon, t_lat, brg, -15.0)
        p_r = offset_point(t_lon, t_lat, brg, 15.0)
        geom = {
            'type': 'LineString',
            'coordinates': [
                [round(p_l[0], 7), round(p_l[1], 7)],
                [round(p_r[0], 7), round(p_r[1], 7)]
            ]
        }
        
        asset_id = f'asset_{next_asset_num:03d}'
        next_asset_num += 1
        
        if ctype.lower() == 'box':
            category = 'Cross Drainage'
            color = '#334155'
            icon = 'culvert_box'
            if cno == 1:
                typology = f'Single Box Culvert 1x({sec}m)'
                short_code = '1x Box Culv'
            elif cno == 2:
                typology = f'Twin Box Culvert 2x({sec}m)'
                short_code = '2x Box Culv'
            elif cno == 3:
                typology = f'Triple Box Culvert 3x({sec}m)'
                short_code = '3x Box Culv'
            else:
                typology = f'{cno}x Box Culvert {cno}x({sec}m)'
                short_code = f'{cno}x Box Culv'
            specs = f'Reinforced concrete box culvert {cno}x({sec}m) [Culvert ID {cid}, Basin {basin}, Q={q} m3/s]'
            drawing_ref = f'{sheet_ref} / MDDTDW220010002 / RP-00001 App II'
        else:
            category = 'Cross Drainage'
            color = '#475569'
            icon = 'culvert_pipe'
            sec_clean = str(sec).strip()
            if not sec_clean.startswith('Ø') and not sec_clean.startswith('O'):
                sec_clean = 'Ø' + sec_clean
            if cno == 1:
                typology = f'Single Pipe Culvert 1x{sec_clean}'
                short_code = 'Pipe Culv'
            else:
                typology = f'{cno}x Pipe Culvert {cno}x{sec_clean}'
                short_code = f'{cno}x Pipe Culv'
            specs = f'Precast concrete pipe culvert {cno}x{sec_clean} [Culvert ID {cid}, Basin {basin}, Q={q} m3/s]'
            drawing_ref = f'{sheet_ref} / MDDTDW220000023-31 / RP-00001 App II'
            
        feature = {
            'type': 'Feature',
            'id': asset_id,
            'properties': {
                'id': asset_id,
                'chainage_str': ch_formatted,
                'start_pk': pk,
                'end_pk': pk,
                'length_m': 0.0,
                'is_point': True,
                'side': 'Center',
                'position': 'Cross Drainage',
                'typology': typology,
                'short_code': short_code,
                'category': category,
                'color': color,
                'drawing_ref': drawing_ref,
                'specs': specs,
                'icon': icon,
                'status': 'Not Started',
                'notes': '',
                'inspection_date': ''
            },
            'geometry': geom
        }
        added_features.append(feature)
        changelog_entries.append(
            f'| {asset_id} | {ch_formatted} | ADD | Culverts (Row {r}) | {drawing_ref} | Added cross-drainage structure {typology} (ID {cid}, Basin {basin}, Q={q} m3/s). Initial field data intact. |'
        )

# Update the 2 VERIFY rows
verify_updates = [
    ('asset_094', 'DW-03009-06', 'Row 53', 'CH 95+212', 'Single Box Culvert 1x(2.5x2.5m)'),
    ('asset_133', 'DW-03020-06', 'Row 54', 'CH 110+111', 'Twin Box Culvert 2x(3.0x3.0m)')
]

for fid, sheet_ref, reg_row, ch_str, typ_str in verify_updates:
    for feat in features:
        if feat['properties']['id'] == fid:
            existing_notes = feat['properties'].get('notes', '')
            flag_text = f'[FLAG] Unsubstantiated by {sheet_ref} and omitted from Design Report RP-00001 Rev 09 Appendix II. Marked for site verification.'
            feat['properties']['notes'] = (existing_notes + ' ' + flag_text).strip()
            changelog_entries.append(
                f'| {fid} | {ch_str} | FLAG | Culverts ({reg_row}) | {sheet_ref} | {flag_text} Existing record retained. |'
            )

# Append new features
features.extend(added_features)
assets_data['metadata']['count'] = len(features)

print(f'Total features after Phase 2: {len(features)} (Added {len(added_features)})')

# Diff against snapshot
for fid, s_data in snapshot.items():
    matching = [f for f in features if f['properties']['id'] == fid]
    assert len(matching) == 1, f'Missing feature {fid}'
    f_props = matching[0]['properties']
    # status must match exactly
    assert f_props['status'] == s_data['status'], f'Status modified for {fid}'
    # if fid not in verify updates, notes must match snapshot exactly
    if fid not in ['asset_094', 'asset_133']:
        assert f_props['notes'] == s_data['notes'], f'Notes modified for {fid}'
    assert f_props['inspection_date'] == s_data['inspection_date'], f'Inspection date modified for {fid}'

print('Field data verification passed! 0 records lost or overwritten.')

# Write back section03_assets.json
with open('data/section03_assets.json', 'w', encoding='utf-8') as f:
    json.dump(assets_data, f, indent=2, ensure_ascii=False)

# Append to CORRECTIONS_CHANGELOG.md
with open('CORRECTIONS_CHANGELOG.md', 'a', encoding='utf-8') as f:
    f.write('\n### Phase 2 — Cross-Drainage Corrections (Culverts Tab)\n\n')
    f.write('| Record ID | Chainage Extent | Action | Register Reference | Drawing / Document Reference | Notes & Field Data Continuity |\n')
    f.write('|:---|:---|:---|:---|:---|:---|\n')
    for line in changelog_entries:
        f.write(line + '\n')

print('CORRECTIONS_CHANGELOG.md updated successfully.')
